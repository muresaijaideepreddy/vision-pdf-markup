"""Synthetic fixtures to exercise plumbing, not measure recognition accuracy."""

from __future__ import annotations

import io
from pathlib import Path


def _word(path: Path, final: bool = False) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    doc.add_paragraph("Markup editing demonstration", "Title")
    doc.add_paragraph("Synthetic sample for testing document editing.")
    p = doc.add_paragraph("The ")
    p.add_run("dog" if final else "cat").bold = True
    p.add_run(" was on the ground.")
    doc.add_paragraph("The device is now ready." if final else "The device is ready.")
    doc.add_paragraph("Status: review." if final else "Status: preliminary review.")
    doc.add_paragraph("The dose was 10 mg daily.")
    table = doc.add_table(rows=2, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Sample A"
    table.cell(1, 1).text = "complete" if final else "pending"
    doc.save(path)


def _scan(path: Path) -> None:
    import pypdfium2 as pdfium
    from reportlab.lib.utils import ImageReader
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.pdfbase.pdfmetrics import stringWidth

    buf = io.BytesIO()
    c = Canvas(buf, pagesize=(612, 792))
    c.setTitle("Synthetic marked scan for a software demonstration")
    c.setFont("Helvetica-Bold", 20)
    c.drawString(56, 730, "Markup editing demonstration")
    c.setFont("Helvetica", 11)
    c.drawString(56, 704, "Synthetic sample for testing document editing.")
    lines = [(650, "The cat was on the ground."), (570, "The device is ready."),
             (490, "Status: preliminary review."), (410, "The dose was 10 mg daily.")]
    for y, text in lines:
        c.drawString(56, y, text)
    c.rect(56, 270, 350, 64)
    c.line(56, 302, 406, 302)
    c.line(231, 270, 231, 334)
    for x, y, text in [(65, 313, "Item"), (240, 313, "Value"), (65, 281, "Sample A"), (240, 281, "pending")]:
        c.drawString(x, y, text)
    c.setStrokeColorRGB(0.08, 0.23, 0.7)
    c.setFillColorRGB(0.08, 0.23, 0.7)
    c.setLineWidth(1.5)
    c.setFont("Times-Italic", 17)
    x = 56 + stringWidth("The ", "Helvetica", 11)
    c.line(x - 1, 654, x + 17, 653)
    c.drawString(x, 671, "dog")
    x = 56 + stringWidth("The device is ", "Helvetica", 11)
    c.line(x - 4, 562, x, 569)
    c.line(x, 569, x + 4, 562)
    c.drawString(x - 2, 587, "now")
    x = 56 + stringWidth("Status: ", "Helvetica", 11)
    c.line(x, 494, x + stringWidth("preliminary ", "Helvetica", 11), 493)
    c.drawString(365, 487, "delete")
    c.line(359, 489, x + 70, 491)
    c.line(239, 285, 280, 286)
    c.drawString(282, 279, "complete")
    c.drawString(247, 410, "? verify")
    c.line(242, 415, 145, 414)
    c.setFillColorRGB(0.35, 0.35, 0.35)
    c.setFont("Helvetica", 9)
    c.drawString(56, 80, "Simulated pen marks. This is not a real handwriting sample.")
    c.drawString(56, 64, "Page 1")
    c.save()
    with pdfium.PdfDocument(buf.getvalue()) as pdf:
        page = pdf[0]
        try:
            bitmap = page.render(scale=2)
            try:
                image = bitmap.to_pil().copy()
            finally:
                bitmap.close()
        finally:
            page.close()
    # Flatten to a raster-only PDF to exercise the scanned-document input path.
    c = Canvas(str(path), pagesize=(612, 792))
    c.setTitle("Synthetic raster marked scan")
    c.drawImage(ImageReader(image), 0, 0, width=612, height=792)
    c.showPage()
    c.save()
    image.close()


def run_demo(destination: Path) -> None:
    from .cli import write_json, write_text
    from .docx_engine import apply_docx, inspect_docx
    from .evaluation import evaluate_documents, write_evaluation
    from .pdf_input import render_pdf_pages
    from .review import render_review

    sample = destination / "sample"
    sample.mkdir()
    original, reference, scan = sample / "original.docx", sample / "reference.docx", sample / "marked-scan.pdf"
    _word(original)
    _word(reference, final=True)
    _scan(scan)
    snapshot = inspect_docx(original)
    snapshot["source"] = str(original.resolve())
    by_text = {block["text"]: block["id"] for block in snapshot["blocks"]}

    def edit(number, text, operation, old, new, before="", after="", review=False, reason="Simulated visible mark"):
        return {"id": f"e{number:03d}", "block_id": by_text[text], "operation": operation, "old_text": old,
                "new_text": new, "context_before": before, "context_after": after, "page": 1,
                "reason": reason, "needs_review": review}

    plan = {
        "schema_version": 1, "source_sha256": snapshot["sha256"],
        "metadata": {"mode": "offline_demo", "scan_pages": 1, "preview_dir": "scan-pages", "model": None},
        "warnings": ["Synthetic offline demo with supplied edits: no model was called and no handwriting accuracy was measured.",
                     "The dose note is unresolved; the demonstration reference retains the original dose.",
                     "Word layout has not been rendered or visually verified by this program."],
        "edits": [
            edit(1, "The cat was on the ground.", "replace", "cat", "dog", "The ", " was", reason="Cross-out cat with replacement dog"),
            edit(2, "The device is ready.", "insert", "", "now ", "The device is ", "ready.", reason="Insert now at caret"),
            edit(3, "Status: preliminary review.", "delete", "preliminary ", "", "Status: ", "review.", reason="Delete preliminary and its following space"),
            edit(4, "pending", "replace", "pending", "complete", reason="Table cell replacement"),
            edit(5, "The dose was 10 mg daily.", "replace", "10", "", "The dose was ", " mg daily.", True, "Question mark requests verification but gives no replacement dose"),
        ],
    }
    render_pdf_pages(scan, destination / "scan-pages")
    write_json(destination / "original-blocks.json", snapshot)
    write_json(destination / "edits.json", plan)
    report = apply_docx(original, plan, destination / "edited.docx")
    write_json(destination / "application-report.json", report)
    write_text(destination / "review.html", render_review(snapshot, plan, report))
    result = evaluate_documents(original, destination / "edited.docx", reference)
    result["experiment_note"] = "OFFLINE DEMO: supplied edits and synthetic reference; not a handwriting-recognition benchmark. One unresolved mark remains."
    write_evaluation(result, destination / "evaluation")
    write_text(destination / "DEMO-NOTES.txt", "Offline demonstration only.\nFour explicit text edits should apply; one question mark should remain for review.\nNo AI model was used. A correct demo result does not measure handwriting recognition.\nThe Word files have not been visually rendered; inspect them in Word before judging layout.\n")
